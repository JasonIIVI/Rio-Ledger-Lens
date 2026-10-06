import hashlib
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

from ledgerlens.cli import build_parser, main


def test_generate_writes_both_files(tmp_path):
    code = main(["generate", "--start", "2024-01-01", "--end", "2024-02-29",
                 "--out-dir", str(tmp_path)])
    assert code == 0
    assert (tmp_path / "ledger.csv").exists()
    assert (tmp_path / "labels.csv").exists()


def test_test_command_runs_against_generated_data(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    code = main(["test", str(tmp_path / "ledger.csv"),
                 "--labels", str(tmp_path / "labels.csv")])
    assert code == 0
    out = capsys.readouterr().out
    assert "Precision" in out
    assert "Recall by archetype" in out


def test_test_command_can_run_a_single_test(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    main(["test", str(tmp_path / "ledger.csv"), "--only", "JET-02"])
    out = capsys.readouterr().out
    assert "Ran 1 test(s)" in out


def test_benford_command_reports_conformity(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2025-12-31",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    main(["benford", str(tmp_path / "ledger.csv")])
    out = capsys.readouterr().out
    assert "MAD" in out
    assert "conformity" in out


def test_bad_date_is_rejected():
    parser = build_parser()
    try:
        parser.parse_args(["generate", "--start", "01/01/2024"])
    except SystemExit as exc:
        assert exc.code != 0
    else:  # pragma: no cover
        raise AssertionError("expected SystemExit")


def test_score_command_reports_tier_agreement(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-12-31",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    code = main(["score", str(tmp_path / "ledger.csv"),
                 "--labels", str(tmp_path / "labels.csv")])
    assert code == 0
    out = capsys.readouterr().out
    assert "Tier agreement" in out
    assert "Isolation Forest" in out
    assert "lift_vs_random" in out


def test_score_command_writes_csv(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    out_csv = tmp_path / "combined.csv"
    main(["score", str(tmp_path / "ledger.csv"), "--out", str(out_csv)])
    assert out_csv.exists()


def test_report_command_writes_workpaper(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-12-31",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    xlsx = tmp_path / "wp.xlsx"
    code = main(["report", str(tmp_path / "ledger.csv"),
                 "--labels", str(tmp_path / "labels.csv"), "--out", str(xlsx)])
    assert code == 0
    assert xlsx.exists()
    assert "Workpaper written" in capsys.readouterr().out


def test_report_command_can_skip_the_model(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    xlsx = tmp_path / "rules_only.xlsx"
    main(["report", str(tmp_path / "ledger.csv"), "--no-model", "--out", str(xlsx)])
    assert xlsx.exists()


def test_narrate_without_a_key_explains_and_fails(tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.chdir(tmp_path)  # a developer's real .env must not leak into the test
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    code = main(["narrate", str(tmp_path / "ledger.csv"), "--db", str(tmp_path / "r.sqlite"),
                 "--no-model"])
    assert code == 1
    assert ".env" in capsys.readouterr().out


def test_narrate_caches_narratives_and_moves_on_next_time(tmp_path, capsys, monkeypatch, llm):
    from ledgerlens import cli
    from ledgerlens.ingest import ledger_identity, load_csv
    from ledgerlens.narrate import Narrator
    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    db = tmp_path / "review.sqlite"
    ledger = str(tmp_path / "ledger.csv")
    ident = ledger_identity(load_csv(ledger), ledger)

    assert main(["narrate", ledger, "--db", str(db), "--top", "4", "--no-model"]) == 0
    out = capsys.readouterr().out
    assert "Narrated 4/4" in out
    assert "read from cache" in out
    assert len(ReviewStore(db, ident).narrative_ids()) == 4
    assert client.calls[0]["output_config"]["effort"] == "medium"

    # Cached entries are skipped, so the next run narrates the next four down.
    assert main(["narrate", ledger, "--db", str(db), "--top", "4", "--no-model"]) == 0
    assert "Skipped 4" in capsys.readouterr().out
    assert len(ReviewStore(db, ident).narrative_ids()) == 8
    assert len(client.calls) == 8

    # --force redoes the top four instead of moving on.
    assert main(["narrate", ledger, "--db", str(db), "--top", "4", "--no-model", "--force"]) == 0
    assert len(ReviewStore(db, ident).narrative_ids()) == 8
    assert len(client.calls) == 12


def test_report_attaches_an_existing_review_db_only(tmp_path, capsys):
    from ledgerlens.ingest import ledger_identity, load_csv
    from ledgerlens.review import Decision, ReviewStore

    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    missing = tmp_path / "absent.sqlite"
    main(["report", str(tmp_path / "ledger.csv"), "--no-model", "--db", str(missing),
          "--out", str(tmp_path / "a.xlsx")])
    assert "without review columns" in capsys.readouterr().out
    assert not missing.exists()  # a report must not create a review database

    db, csv = tmp_path / "review.sqlite", str(tmp_path / "ledger.csv")
    ident = ledger_identity(load_csv(csv), csv)
    ReviewStore(db, ident).record(Decision("JE-2024-000001", "dismiss", "ana"))
    other = ReviewStore(db, "csv:" + "b" * 64)  # another ledger's decisions, same file
    for entry in ("JE-2024-000001", "JE-2024-000002"):
        other.record(Decision(entry, "escalate", "bo"))
    code = main(["report", str(tmp_path / "ledger.csv"), "--no-model", "--db", str(db),
                 "--out", str(tmp_path / "b.xlsx")])
    assert code == 0
    # The workpaper's review section is this ledger's, and the other ledger is only counted.
    import openpyxl

    sheet = openpyxl.load_workbook(tmp_path / "b.xlsx")["Summary"]
    summary = {row[0].value: row[1].value for row in sheet.iter_rows() if row[0].value}
    assert summary["Review rows for ledger"] == ident
    assert summary["Decisions recorded"] == 1
    assert summary["  dismiss"] == 1 and "  escalate" not in summary
    assert summary["Other ledgers in this database (not shown)"] == (
        "1 ledger(s): 0 narrated, 2 decided entries")


def test_eval_narratives_select_writes_a_skeleton_and_will_not_clobber_it(tmp_path, capsys):
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    ledger, labels = str(tmp_path / "ledger.csv"), str(tmp_path / "labels.csv")
    cases = tmp_path / "cases.json"

    assert main(["eval-narratives", ledger, "--select", "--cases", str(cases)]) == 2  # no labels
    assert main(["eval-narratives", ledger, "--select", "--labels", labels,
                 "--cases", str(cases)]) == 0
    payload = json.loads(cases.read_text())
    assert payload["cases"] and all("must_mention" in c for c in payload["cases"])
    assert main(["eval-narratives", ledger, "--select", "--labels", labels,
                 "--cases", str(cases)]) == 2  # exists, no --overwrite


def test_eval_narratives_defaults_to_a_dated_runs_dir_and_regrades_the_newest(
        tmp_path, capsys, monkeypatch, llm):
    from ledgerlens import cli
    from ledgerlens.ingest import load_csv
    from ledgerlens.narrate import DEFAULT_MODEL, Narrator

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    ledger, labels = str(tmp_path / "ledger.csv"), str(tmp_path / "labels.csv")
    # The default runs directory is for the default ledger only. Declaring this
    # six-month ledger to be it keeps the test off the 10,000-line one (the pin
    # test covers the real constant, CSV round trip included).
    monkeypatch.setattr(cli.narrative_eval, "DEFAULT_LEDGER_SHA256",
                        cli.narrative_eval.ledger_digest(load_csv(ledger)))
    cases, report = tmp_path / "cases.json", tmp_path / "report.md"
    main(["eval-narratives", ledger, "--select", "--labels", labels, "--cases", str(cases)])
    capsys.readouterr()
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    argv = ["eval-narratives", ledger, "--cases", str(cases), "--out", str(report), "--limit", "2"]

    # Nothing to re-grade yet is an error that names the flag, not a fresh paid run,
    # whether the directory is defaulted or mistyped; and a re-grade cannot start over.
    assert main(argv + ["--regrade"]) == 2
    assert "runs-dir" in capsys.readouterr().out
    assert main(argv + ["--regrade", "--runs-dir", str(tmp_path / "typo")]) == 2
    assert "nothing to re-grade" in capsys.readouterr().out
    with pytest.raises(SystemExit) as refused:
        main(argv + ["--regrade", "--no-resume"])
    assert refused.value.code == 2
    assert client.calls == []

    assert main(argv) == 0
    runs = sorted((tmp_path / "evals" / "narratives" / "runs").iterdir())
    assert len(runs) == 1 and runs[0].name.endswith(f"-{DEFAULT_MODEL}")
    assert len((runs[0] / "results.jsonl").read_text().splitlines()) == 2
    assert "evals/narratives/runs/" in report.read_text()

    # A re-grade finds that directory without being told, and calls nothing.
    assert main(argv + ["--regrade"]) == 0
    assert len(client.calls) == 2
    assert "re-graded" in report.read_text()

    # Only the expected refusals become exit 2; anything else is a bug and surfaces.
    def boom(*args, **kwargs):
        raise ValueError("boom")
    monkeypatch.setattr(cli.narrative_eval, "run_eval", boom)
    with pytest.raises(ValueError, match="boom"):
        main(argv + ["--regrade"])


def test_eval_narratives_keeps_other_ledgers_out_of_the_committed_paths(tmp_path, capsys,
                                                                       monkeypatch, llm):
    """Rule 1 at the point of writing: another ledger needs its own --runs-dir and --cases."""
    from ledgerlens import cli
    from ledgerlens.ingest import load_csv
    from ledgerlens.narrate import Narrator

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    ledger, labels = str(tmp_path / "ledger.csv"), str(tmp_path / "labels.csv")
    capsys.readouterr()

    assert main(["eval-narratives", ledger, "--select", "--labels", labels]) == 2  # default --cases
    assert "default synthetic ledger" in capsys.readouterr().out
    assert not (tmp_path / "evals").exists()
    cases = tmp_path / "cases.json"
    assert main(["eval-narratives", ledger, "--select", "--labels", labels,
                 "--cases", str(cases)]) == 0
    capsys.readouterr()

    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    argv = ["eval-narratives", ledger, "--cases", str(cases), "--out", str(tmp_path / "r.md"),
            "--limit", "2"]
    assert main(argv) == 2  # default --runs-dir would be the committed directory
    out = capsys.readouterr().out
    assert "default synthetic ledger" in out and "runs" in out
    assert client.calls == [] and not (tmp_path / "evals").exists()
    # Nor does spelling a committed path out, an empty one, or the report's default.
    for spelled in (["--runs-dir", "evals/narratives/runs/2026-10-01-claude-x"],
                    ["--runs-dir", "./evals/narratives/runs/x"], ["--runs-dir", "docs"]):
        assert main(argv + spelled) == 2, spelled
        assert "committed" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        main(argv + ["--runs-dir", ""])
    capsys.readouterr()
    no_out = [a for a in argv if a not in ("--out", str(tmp_path / "r.md"))]
    assert main(no_out + ["--runs-dir", str(tmp_path / "runs")]) == 2  # docs/narrative-eval.md
    assert "the report" in capsys.readouterr().out
    assert main(["eval-narratives", ledger, "--select", "--labels", labels,
                 "--cases", "./evals/narratives/cases.json", "--overwrite"]) == 2
    capsys.readouterr()
    assert client.calls == []
    assert not (tmp_path / "evals").exists() and not (tmp_path / "docs").exists()

    assert main(argv + ["--runs-dir", str(tmp_path / "runs")]) == 0
    out = capsys.readouterr().out
    assert "not the default synthetic ledger" in out
    rows = [json.loads(line) for line in (tmp_path / "runs" / "results.jsonl").read_text().splitlines()]
    digest = cli.narrative_eval.ledger_digest(load_csv(ledger))
    assert digest != cli.narrative_eval.DEFAULT_LEDGER_SHA256
    assert len(rows) == 2 and all(r["ledger_sha256"] == digest for r in rows)
    assert "not the default synthetic ledger" in (tmp_path / "r.md").read_text()


def test_eval_narratives_names_a_missing_or_damaged_case_file(tmp_path, capsys):
    """Run from a directory that is not the checkout, the default case path is missing: a
    message and exit 2, not a traceback; a file that is not a case set likewise."""
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    capsys.readouterr()
    common = ["eval-narratives", str(tmp_path / "ledger.csv"),
              "--runs-dir", str(tmp_path / "runs"), "--out", str(tmp_path / "report.md")]
    assert main(common + ["--cases", str(tmp_path / "missing.json")]) == 2
    assert "case file not found" in capsys.readouterr().out
    (tmp_path / "damaged.json").write_text("{not json")
    assert main(common + ["--cases", str(tmp_path / "damaged.json")]) == 2
    assert "damaged.json" in capsys.readouterr().out

    # Every case the loader refuses is a message too, whatever the reason.
    good = {"entry_id": "JE-2024-000001", "archetype": "a", "tests_fired": "JET-01",
            "why_chosen": "w"}
    refused = {
        "a pattern that is not a regex": {"cases": [{**good, "must_mention": ["1,?322("]}]},
        "a misspelt key": {"cases": [{**good, "overides": {}}]},
        "a missing field": {"cases": [{"entry_id": "JE-2024-000001"}]},
        "no case list": {"about": "x"},
        "a list, not an object": [good],
        "overrides that are not an object": {"cases": [{**good, "overrides": 5}]},
        "overrides that are null": {"cases": [{**good, "overrides": None}]},
        "a pattern list that is a string": {"cases": [{**good, "must_mention": "1,322"}]},
    }
    for why, payload in refused.items():
        (tmp_path / "refused.json").write_text(json.dumps(payload))
        assert main(common + ["--cases", str(tmp_path / "refused.json")]) == 2, why
        out = capsys.readouterr().out
        assert out.startswith("error:") and "refused.json" in out, why
    assert main(common + ["--cases", str(tmp_path)]) == 2  # a directory
    assert capsys.readouterr().out.startswith("error:")


def test_eval_narratives_checks_override_lines_before_narrating_anything(
        tmp_path, capsys, monkeypatch, llm):
    """An override on a line the entry lacks is a stale case set, not a traceback mid-run."""
    from ledgerlens import cli
    from ledgerlens.ingest import load_csv
    from ledgerlens.narrate import Narrator

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30", "--out-dir", str(tmp_path)])
    ledger, labels = str(tmp_path / "ledger.csv"), str(tmp_path / "labels.csv")
    cases = tmp_path / "cases.json"
    main(["eval-narratives", ledger, "--select", "--labels", labels, "--cases", str(cases)])
    payload = json.loads(cases.read_text())
    last = payload["cases"][-1]
    beyond = int(load_csv(ledger).query("entry_id == @last['entry_id']")["line_no"].max()) + 1
    last["overrides"] = {"lines": {str(beyond): {"description": "an extra line"}}}
    cases.write_text(json.dumps(payload))
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    capsys.readouterr()

    assert main(["eval-narratives", ledger, "--cases", str(cases), "--out", str(tmp_path / "r.md"),
                 "--runs-dir", str(tmp_path / "runs")]) == 2
    out = capsys.readouterr().out
    assert f"{last['entry_id']}: no line {beyond} to override" in out
    assert client.calls == [] and not (tmp_path / "runs" / "results.jsonl").exists()


def test_eval_narratives_rejects_a_limit_below_one(capsys):
    """--limit 0 must not mean "every case", which is a paid run."""
    with pytest.raises(SystemExit) as refused:
        main(["eval-narratives", "x.csv", "--limit", "0"])
    assert refused.value.code == 2
    assert "at least 1" in capsys.readouterr().err


def test_eval_narratives_grades_the_cases_and_writes_the_report(tmp_path, capsys, monkeypatch,
                                                                 llm):
    from ledgerlens import cli
    from ledgerlens.narrate import Narrator

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-06-30",
          "--out-dir", str(tmp_path)])
    ledger, labels = str(tmp_path / "ledger.csv"), str(tmp_path / "labels.csv")
    cases = tmp_path / "cases.json"
    main(["eval-narratives", ledger, "--select", "--labels", labels, "--cases", str(cases)])
    capsys.readouterr()
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))

    report, runs = tmp_path / "report.md", tmp_path / "runs"
    argv = ["eval-narratives", ledger, "--cases", str(cases), "--out", str(report),
            "--runs-dir", str(runs), "--limit", "3"]
    code = main(argv)
    out = capsys.readouterr().out
    assert code == 0, out
    assert "Graded 3/3" in out
    assert report.read_text().startswith("# Narrative eval")
    rows = [json.loads(line) for line in (runs / "results.jsonl").read_text().splitlines()]
    assert len(rows) == 3
    assert len(client.calls) == 3

    # Every row and the report say which case file graded them; the report also
    # carries the grader's history and the constant-answer confidence baseline.
    digest = hashlib.sha256(cases.read_bytes()).hexdigest()
    assert all(r["cases_sha256"] == digest for r in rows)
    text = report.read_text()
    assert digest in text and "## Grader notes" in text and "Confidence baseline" in text

    # Editing a band after the run: a plain re-grade is refused, an explicit one is disclosed.
    payload = json.loads(cases.read_text())
    payload["cases"][0]["expected_confidence"] = ["low"]
    cases.write_text(json.dumps(payload))
    assert main(argv + ["--regrade"]) == 2
    assert "allow-cases-change" in capsys.readouterr().out
    assert main(argv + ["--regrade", "--allow-cases-change"]) == 0
    assert len(client.calls) == 3  # neither re-grade called the API
    text = report.read_text()
    assert "Provenance note" in text and digest in text
    assert hashlib.sha256(cases.read_bytes()).hexdigest() in text

    # Starting over into a directory that holds rows is refused, not overwritten.
    assert main(argv + ["--no-resume"]) == 2
    assert "never deleted" in capsys.readouterr().out
    assert len(client.calls) == 3


def test_narrate_and_report_refuse_a_database_they_cannot_open_without_a_traceback(
        tmp_path, capsys, monkeypatch):
    import sqlite3

    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31", "--out-dir", str(tmp_path)])
    ledger, out = str(tmp_path / "ledger.csv"), str(tmp_path / "w.xlsx")
    future = tmp_path / "future.sqlite"
    ReviewStore(future, "csv:" + "a" * 64)
    with sqlite3.connect(str(future)) as raw:
        raw.execute("PRAGMA user_version = 99")
    capsys.readouterr()

    assert main(["report", ledger, "--no-model", "--db", str(future), "--out", out]) == 2
    assert "newer LedgerLens" in capsys.readouterr().out
    # narrate opens the store before it builds an API client, so this needs no key.
    assert main(["narrate", ledger, "--no-model", "--db", str(future)]) == 2
    assert "newer LedgerLens" in capsys.readouterr().out
    assert main(["adopt-legacy", ledger, "--db", str(future)]) == 2
    assert "newer LedgerLens" in capsys.readouterr().out

    notes = tmp_path / "notes.sqlite"
    notes.write_text("not a database\n")
    assert main(["report", ledger, "--no-model", "--db", str(notes), "--out", out]) == 2
    assert "not a SQLite database" in capsys.readouterr().out
    assert notes.read_text() == "not a database\n"


def test_narrate_files_each_ledger_under_its_own_identity(tmp_path, capsys, monkeypatch, llm):
    """Two ledgers that share entry ids share one database and never each other's notes."""
    from ledgerlens import cli
    from ledgerlens.ingest import ledger_identity, load_csv
    from ledgerlens.narrate import Narrator
    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    db = tmp_path / "review.sqlite"
    ledgers = []
    for end in ("2024-03-31", "2024-06-30"):
        out = tmp_path / end
        main(["generate", "--start", "2024-01-01", "--end", end, "--out-dir", str(out)])
        ledgers.append(str(out / "ledger.csv"))
    capsys.readouterr()
    for ledger in ledgers:
        assert main(["narrate", ledger, "--db", str(db), "--top", "4", "--no-model"]) == 0
        out = capsys.readouterr().out
        assert "Skipped" not in out  # the second ledger's run starts from nothing of its own
    identities = [ledger_identity(load_csv(path), path) for path in ledgers]
    assert identities[0] != identities[1]
    for ident in identities:
        assert len(ReviewStore(db, ident).narrative_ids()) == 4
    assert ReviewStore(db, identities[0]).ledgers()["ledger_id"].tolist() == sorted(identities)
    assert len(client.calls) == 8


def test_narrate_refuses_to_ignore_legacy_rows_unless_told_to(tmp_path, capsys, monkeypatch, llm):
    """Rows from before ledgers were keyed are adopted or set aside on purpose, never bought again."""
    import sqlite3

    from ledgerlens import cli
    from ledgerlens.ingest import ledger_identity, load_csv
    from ledgerlens.narrate import Narrator
    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31", "--out-dir", str(tmp_path)])
    ledger, db = str(tmp_path / "ledger.csv"), tmp_path / "review.sqlite"
    ReviewStore(db, ledger_identity(load_csv(ledger), ledger))
    with sqlite3.connect(str(db)) as raw:  # a note from before ledgers were keyed
        raw.execute("INSERT INTO narratives (entry_id, summary, generated_at, ledger_id) VALUES "
                    "('JE-2024-000001', 'old', '2026-09-23T00:00:00+00:00', 'legacy')")
    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    capsys.readouterr()

    # The refusal comes before anything is narrated: asked for four notes, it buys none.
    assert main(["narrate", ledger, "--db", str(db), "--top", "4", "--no-model"]) == 2
    out = capsys.readouterr().out
    assert "adopt-legacy" in out and client.calls == []
    assert ReviewStore(db, ledger_identity(load_csv(ledger), ledger)).narrative_ids() == set()
    fresh = tmp_path / "fresh.sqlite"
    ReviewStore(fresh, ledger_identity(load_csv(ledger), ledger))
    assert main(["adopt-legacy", ledger, "--db", str(fresh)]) == 0  # no legacy rows at all
    assert "Nothing to adopt" in capsys.readouterr().out
    assert main(["adopt-legacy", ledger, "--db", str(tmp_path / "absent.sqlite")]) == 2
    assert main(["adopt-legacy", ledger, "--db", str(db)]) == 0
    assert "Adopted 1 narrative(s) and 0 decision(s)" in capsys.readouterr().out
    assert main(["adopt-legacy", ledger, "--db", str(db)]) == 2
    assert "already holds" in capsys.readouterr().out
    assert main(["narrate", ledger, "--db", str(db), "--top", "0", "--no-model"]) == 0
    assert "Skipped 1" in capsys.readouterr().out

    # Rows for entries this ledger does not have are refused; the ledger may
    # then start from scratch when told so.
    other = tmp_path / "other.sqlite"
    ReviewStore(other, ledger_identity(load_csv(ledger), ledger))
    with sqlite3.connect(str(other)) as raw:
        raw.execute("INSERT INTO narratives (entry_id, summary, generated_at, ledger_id) VALUES "
                    "('JE-9999-000001', 'stray', '2026-09-23T00:00:00+00:00', 'legacy')")
    assert main(["adopt-legacy", ledger, "--db", str(other)]) == 2
    assert "not in this ledger" in capsys.readouterr().out
    assert main(["narrate", ledger, "--db", str(other), "--top", "1", "--no-model",
                 "--ignore-legacy"]) == 0
    assert len(client.calls) == 1  # told to start from scratch, it narrates
    assert len(ReviewStore(other, ledger_identity(load_csv(ledger), ledger)).narrative_ids()) == 1


@pytest.mark.parametrize("payload", ['{"ledger_id": "qbo 123"}', ""])
def test_a_malformed_identity_sidecar_stops_every_command_with_a_message(
        tmp_path, capsys, monkeypatch, payload):
    """A sidecar that names no ledger is never guessed around, and never a traceback."""
    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31", "--out-dir", str(tmp_path)])
    ledger, db = str(tmp_path / "ledger.csv"), tmp_path / "review.sqlite"
    ReviewStore(db, "csv:" + "a" * 64)
    (tmp_path / "ledger.identity.json").write_text(payload)
    capsys.readouterr()
    for argv in (["narrate", ledger, "--no-model", "--top", "0", "--db", str(db)],
                 ["report", ledger, "--no-model", "--db", str(db), "--out", str(tmp_path / "w.xlsx")],
                 ["adopt-legacy", ledger, "--db", str(db)]):
        assert main(argv) == 2, argv[0]
        out = capsys.readouterr().out
        assert "error:" in out and "ledger.identity.json" in out, argv[0]


# --- QuickBooks Online: pull-qbo and qbo-auth (offline) ---------------------------

QBO_FIXTURES = Path(__file__).resolve().parent / "fixtures" / "qbo"
QBO_ENV = {"QBO_CLIENT_ID": "CLIENT-ID", "QBO_CLIENT_SECRET": "CLIENT-SECRET",
           "QBO_ENVIRONMENT": "sandbox"}
QBO_PERIOD = ["--start", "2026-07-01", "--end", "2026-09-30"]  # the recorded sandbox quarter


@pytest.fixture
def qbo_env(tmp_path, monkeypatch):
    """A clean QuickBooks environment: no .env in the working directory, no QBO_* leaking in
    from the shell, and a token directory outside any git checkout."""
    from ledgerlens.connectors.tokens import repository_root

    if repository_root(tmp_path) is not None:
        pytest.skip(f"{tmp_path} is inside a git repository")
    monkeypatch.chdir(tmp_path)
    for name in ("QBO_CLIENT_ID", "QBO_CLIENT_SECRET", "QBO_ENVIRONMENT", "QBO_REALM_ID",
                 "QBO_REDIRECT_URI", "QBO_TIMEZONE", "LEDGERLENS_TOKEN_DIR"):
        monkeypatch.delenv(name, raising=False)
    tokens = tmp_path / "tokens"
    monkeypatch.setenv("LEDGERLENS_TOKEN_DIR", str(tokens))
    return SimpleNamespace(root=tmp_path, tokens=tokens, set=lambda **kv: [
        monkeypatch.setenv(k, v) for k, v in {**QBO_ENV, **kv}.items()])


def test_pull_qbo_from_fixtures_writes_a_ledger_every_command_reads(qbo_env, capsys):
    out = qbo_env.root / "data" / "qbo-ledger.csv"
    code = main(["pull-qbo", *QBO_PERIOD, "--fixtures", str(QBO_FIXTURES / "pull"), "--out", str(out)])
    printed = capsys.readouterr().out
    assert code == 0, printed
    assert "Wrote 297 lines / 116 entries" in printed
    assert "Identity qbo:sandbox-fixtures written to" in printed
    assert "small population" not in printed and "not used" not in printed  # 116 entries
    sidecar = out.with_suffix(".identity.json")
    assert json.loads(sidecar.read_text())["ledger_id"] == "qbo:sandbox-fixtures"
    assert main(["test", str(out)]) == 0
    # the sandbox's round $25,000 opening-balance journal entry: a round amount (JET-01) keyed
    # by hand by a user not on the generator's approver list (JET-12)
    top = capsys.readouterr().out.split("Top 10 by risk score:")[1]
    assert re.search(r"QBO-JournalEntry-8 .*JET-01.*JET-12.*\$25,000\.00", top)
    assert main(["report", str(out), "--no-model", "--out", str(qbo_env.root / "wp.xlsx")]) == 0


def test_a_pull_below_the_small_population_line_says_the_tiers_need_more(qbo_env, capsys, monkeypatch):
    from ledgerlens.connectors import qbo

    monkeypatch.setattr(qbo, "SMALL_LEDGER", 117)  # the recorded quarter holds 116 entries
    assert main(["pull-qbo", *QBO_PERIOD, "--fixtures", str(QBO_FIXTURES / "pull"),
                 "--out", str(qbo_env.root / "data" / "q.csv")]) == 0
    assert "Warning: 116 entries is a small population" in capsys.readouterr().out


@pytest.mark.parametrize("argv, env, expected", [
    (["--start", "2025-12-31", "--end", "2025-10-01"], {}, "is after --end"),
    ([], None, "missing QBO_CLIENT_ID, QBO_CLIENT_SECRET"),
    ([], {"QBO_ENVIRONMENT": "production", "QBO_REALM_ID": "1"}, "--allow-production"),
    (["--record", "rec"], {"QBO_ENVIRONMENT": "production", "QBO_REALM_ID": "1"}, "only sandbox data"),
    ([], {}, "no realm id"),
    ([], {"QBO_REALM_ID": "4620816365"}, "no tokens for sandbox company 4620816365.*qbo-auth"),
    (["--out", "data/ledger.txt"], {}, "must be a .csv file"),
    (["--realm-id", "_x"], {}, "not letters, digits"),
    ([], {"QBO_REALM_ID": "12 34"}, "not letters, digits"),
])
def test_pull_qbo_refuses_what_it_cannot_do_with_a_message(qbo_env, capsys, argv, env, expected):
    if env is not None:
        qbo_env.set(**env)
    args = ["pull-qbo", "--start", "2025-10-01", "--end", "2025-12-31", *argv]
    if "--start" in argv:
        args = ["pull-qbo", *argv]
    assert main(args) == 2
    assert re.search(expected, capsys.readouterr().out)
    assert not (qbo_env.root / "data" / "qbo-ledger.csv").exists()


def _stored_tokens(qbo_env, realm="4620816365"):
    from ledgerlens.connectors.tokens import Tokens, TokenStore

    store = TokenStore.for_realm("sandbox", realm, qbo_env.tokens)
    store.save(Tokens.issued("STORED-ACCESS", "STORED-REFRESH", 3600, 86400, realm, "sandbox"))
    return store


def test_a_live_pull_uses_the_stored_tokens_and_records_sanitized_fixtures(qbo_env, capsys, monkeypatch):
    """The live path end to end, with the network replaced by the pull fixtures: the realm
    comes from the only token file, and --record writes fixtures that replay the same pull."""
    from ledgerlens.connectors import qbo

    qbo_env.set()
    _stored_tokens(qbo_env)
    live = qbo.RecordedTransport(QBO_FIXTURES / "pull")
    monkeypatch.setattr(qbo, "default_transport", lambda: live)
    recorded = qbo_env.root / "recorded"
    code = main(["pull-qbo", *QBO_PERIOD, "--record", str(recorded)])
    printed = capsys.readouterr().out
    assert code == 0, printed
    assert "Identity qbo:4620816365" in printed and "Recorded 3 sanitized fixture(s)" in printed
    assert "STORED-ACCESS" not in "".join(p.read_text() for p in recorded.glob("*.json"))
    first = (qbo_env.root / "data" / "qbo-ledger.csv").read_text()
    (qbo_env.root / "data" / "qbo-ledger.csv").unlink()
    assert main(["pull-qbo", *QBO_PERIOD, "--realm-id", "4620816365", "--fixtures", str(recorded)]) == 0
    assert (qbo_env.root / "data" / "qbo-ledger.csv").read_text() == first
    assert "not used" not in capsys.readouterr().out
    assert main(["pull-qbo", *QBO_PERIOD, "--record", str(recorded)]) == 2
    assert "already holds fixtures" in capsys.readouterr().out


def _free_port() -> int:
    import socket

    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        return probe.getsockname()[1]


def test_qbo_auth_signs_in_and_stores_the_tokens_privately(qbo_env, capsys, monkeypatch):
    import stat
    import threading
    import urllib.parse
    import urllib.request
    import webbrowser

    from ledgerlens.connectors import qbo

    port = _free_port()
    qbo_env.set(QBO_REDIRECT_URI=f"http://localhost:{port}/callback")
    monkeypatch.setattr(qbo, "default_transport",
                        lambda: qbo.RecordedTransport(QBO_FIXTURES / "auth"))
    visited = []

    def browser(url):  # the user signs in; Intuit redirects back with a code
        state = dict(urllib.parse.parse_qsl(url.split("?", 1)[1]))["state"]
        back = f"http://127.0.0.1:{port}/callback?" + urllib.parse.urlencode(
            {"code": "ONE-TIME-CODE", "state": state, "realmId": "4620816365"})
        direct = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        thread = threading.Thread(target=lambda: visited.append(direct.open(back, timeout=5).status))
        thread.start()
        threads.append(thread)
        return True

    threads = []
    monkeypatch.setattr(webbrowser, "open", browser)
    code = main(["qbo-auth", "--timeout", "10"])
    for thread in threads:
        thread.join(5)
    printed = capsys.readouterr().out
    assert code == 0, printed
    token_file = qbo_env.tokens / "qbo-sandbox-4620816365.json"
    assert stat.S_IMODE(token_file.stat().st_mode) == 0o600
    assert json.loads(token_file.read_text())["access_token"] == "TEST-ACCESS"
    assert "Add to .env: QBO_REALM_ID=4620816365" in printed
    for secret in ("TEST-ACCESS", "TEST-REFRESH", "ONE-TIME-CODE", "CLIENT-SECRET"):
        assert secret not in printed
    assert visited == [200]


def test_qbo_auth_times_out_without_writing_anything(qbo_env, capsys):
    qbo_env.set(QBO_REDIRECT_URI=f"http://localhost:{_free_port()}/callback")
    assert main(["qbo-auth", "--timeout", "0.3", "--no-browser"]) == 1
    assert "no sign-in arrived" in capsys.readouterr().out
    assert list(qbo_env.tokens.iterdir()) == []  # the private directory is made first; no tokens


def test_a_qbo_sidecar_keys_adopt_legacy_narrate_and_the_report(tmp_path, capsys, monkeypatch, llm):
    """A QuickBooks pull's ledger is filed under qbo:<realm> on every surface that writes or
    reads review rows, whatever the CSV's digest (a re-pull changes the digest, not the key)."""
    import sqlite3

    import openpyxl

    from ledgerlens import cli
    from ledgerlens.narrate import Narrator
    from ledgerlens.review import ReviewStore

    monkeypatch.chdir(tmp_path)
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31", "--out-dir", str(tmp_path)])
    ledger, db = str(tmp_path / "ledger.csv"), tmp_path / "review.sqlite"
    (tmp_path / "ledger.identity.json").write_text('{"ledger_id": "qbo:4620816365"}')
    ReviewStore(db, "qbo:4620816365")
    with sqlite3.connect(str(db)) as raw:
        raw.execute("INSERT INTO narratives (entry_id, summary, generated_at, ledger_id) VALUES "
                    "('JE-2024-000001', 'old', '2026-09-23T00:00:00+00:00', 'legacy')")
    capsys.readouterr()

    assert main(["adopt-legacy", ledger, "--db", str(db)]) == 0
    assert "into qbo:4620816365" in capsys.readouterr().out

    client = llm.client()
    monkeypatch.setattr(cli, "Narrator", lambda **kw: Narrator(client=client, **kw))
    assert main(["narrate", ledger, "--db", str(db), "--top", "2", "--no-model"]) == 0
    assert "rows keyed by qbo:4620816365" in capsys.readouterr().out
    store = ReviewStore(db, "qbo:4620816365")
    assert len(store.narrative_ids()) >= 2 and "JE-2024-000001" in store.narrative_ids()
    assert set(store.ledgers()["ledger_id"]) == {"legacy", "qbo:4620816365"}

    out = tmp_path / "wp.xlsx"
    assert main(["report", ledger, "--no-model", "--db", str(db), "--out", str(out)]) == 0
    cells = {str(c.value) for ws in openpyxl.load_workbook(out) for row in ws.iter_rows()
             for c in row if c.value is not None}
    assert "qbo:4620816365" in cells


def test_nothing_is_spent_when_the_tokens_could_not_be_saved(qbo_env, capsys, monkeypatch):
    """A code is single-use and a refresh rotates the token: a token directory the store
    would refuse is refused before the browser opens or any request is sent."""
    import webbrowser

    from ledgerlens.connectors import qbo

    qbo_env.set(QBO_REDIRECT_URI=f"http://localhost:{_free_port()}/callback")
    store = _stored_tokens(qbo_env)
    qbo_env.tokens.chmod(0o755)
    sent = []
    monkeypatch.setattr(qbo, "default_transport",
                        lambda: SimpleNamespace(request=lambda *a, **k: sent.append(a)))
    monkeypatch.setattr(webbrowser, "open", lambda url: sent.append(url))
    assert main(["qbo-auth", "--timeout", "0.3"]) == 2
    assert "mode is 755" in capsys.readouterr().out
    assert main(["pull-qbo", "--start", "2025-10-01", "--end", "2025-12-31"]) == 2
    assert "mode is 755" in capsys.readouterr().out
    assert sent == [] and store.load().refresh_token == "STORED-REFRESH"
    qbo_env.tokens.chmod(0o700)
    (qbo_env.root / "rec").write_text("a file")
    assert main(["pull-qbo", "--start", "2025-10-01", "--end", "2025-12-31",
                 "--record", str(qbo_env.root / "rec")]) == 2
    assert "is a file" in capsys.readouterr().out and sent == []


def _replay_to(out, *extra):
    return main(["pull-qbo", *QBO_PERIOD, "--fixtures", str(QBO_FIXTURES / "pull"), "--out", str(out), *extra])


def test_the_synthetic_ledger_and_its_identity_are_refused_however_spelled(qbo_env, capsys, monkeypatch):
    (qbo_env.root / ".git").mkdir()  # the working directory is a checkout, as in real use
    data = qbo_env.root / "data"
    data.mkdir()
    synthetic = data / "ledger.csv"
    synthetic.write_text("the synthetic ledger\\n")
    spellings = ["data/ledger.csv", "./data/../data/ledger.csv"]
    if (qbo_env.root / "DATA" / "LEDGER.CSV").exists():  # a case-insensitive disk, as on macOS
        spellings.append("data/LEDGER.csv")
    for spelled in spellings:
        assert _replay_to(spelled) == 2, spelled
        assert "synthetic ledger" in capsys.readouterr().out
    (qbo_env.root / "src").mkdir()
    monkeypatch.chdir(qbo_env.root / "src")
    assert _replay_to("../data/ledger.csv") == 2  # from another directory of the checkout
    monkeypatch.chdir(qbo_env.root)
    assert _replay_to("data/ledger.txt") == 2  # would take over data/ledger.identity.json
    assert synthetic.read_text() == "the synthetic ledger\\n"
    assert not (data / "ledger.identity.json").exists()


def test_pulled_books_stay_in_the_checkouts_ignored_data_directory(qbo_env, capsys):
    repo = qbo_env.root / "repo"
    (repo / ".git").mkdir(parents=True)
    (repo / "tests").mkdir()
    assert _replay_to(repo / "tests" / "q.csv") == 2
    assert "not in its data/ directory" in capsys.readouterr().out
    assert _replay_to(repo / "q.csv") == 2
    assert _replay_to(repo / "data" / "q.csv") == 0
    assert (repo / "data" / "q.identity.json").exists()
    assert _replay_to(qbo_env.root / "elsewhere" / "q.csv") == 0  # outside any checkout


def test_a_replay_converts_times_with_qbo_timezone_like_the_recorded_pull(qbo_env, monkeypatch):
    import pandas as pd

    monkeypatch.setenv("QBO_TIMEZONE", "America/New_York")
    assert _replay_to(qbo_env.root / "data" / "q.csv") == 0
    ledger = pd.read_csv(qbo_env.root / "data" / "q.csv")
    # the recording's -07:00 (entity) and -0700 (report), three hours later in New York
    times = set(ledger.loc[ledger["entry_id"] == "QBO-JournalEntry-6", "entered_at"])
    assert times == {"2026-08-31 15:11:06"}
    times = set(ledger.loc[ledger["entry_id"] == "QBO-Check-57", "entered_at"])
    assert times == {"2026-09-02 18:14:27"}
    monkeypatch.setenv("QBO_TIMEZONE", "Mars/Olympus")
    assert _replay_to(qbo_env.root / "data" / "q.csv") == 2


def test_a_recording_that_ends_in_an_error_is_still_scrubbed_and_checked(qbo_env, capsys, monkeypatch):
    """A name the report reveals last must leave the earlier fixtures even when the pull then
    fails (here: a report without a credit column, as a multicurrency company's would be)."""
    import shutil

    from ledgerlens.connectors import qbo

    live_dir = qbo_env.root / "live"
    shutil.copytree(QBO_FIXTURES / "pull", live_dir)
    je = live_dir / "020-post-query-journalentry-p1.json"
    je.write_text(je.read_text().replace('"PrivateNote": "Opening Balance"',
                                         '"PrivateNote": "Opening balance approved by Jane Dev"'))
    gl = live_dir / "030-get-reports-generalledger.json"
    gl.write_text(gl.read_text().replace('"qbo-user-1"', '"Jane Dev"')
                  .replace('"Value": "credit_amt"', '"Value": "credit_x"')
                  .replace('"ColTitle": "Credit"', '"ColTitle": "Credit X"'))
    qbo_env.set()
    _stored_tokens(qbo_env)
    monkeypatch.setattr(qbo, "default_transport", lambda: qbo.RecordedTransport(live_dir))
    recorded = qbo_env.root / "recorded"
    assert main(["pull-qbo", *QBO_PERIOD, "--record", str(recorded)]) == 1
    printed = capsys.readouterr().out
    assert "multicurrency" in printed and "not a complete pull" in printed
    text = "".join(p.read_text() for p in recorded.glob("*.json"))
    assert "Jane Dev" not in text and "approved by qbo-user-" in text


def test_pulled_books_never_land_in_a_new_directory_of_the_checkout(qbo_env, capsys):
    repo = qbo_env.root / "repo"
    (repo / ".git").mkdir(parents=True)
    for place in (repo / "tests" / "new" / "q.csv", repo / "q2" / "q.csv"):
        assert _replay_to(place) == 2, place
        assert "not in its data/ directory" in capsys.readouterr().out
        assert not place.parent.exists()
    assert _replay_to(qbo_env.root / "data" / "ledger.CSV") == 2  # a suffix in another case
    assert "must be a .csv file" in capsys.readouterr().out


def test_qbo_auth_refuses_a_token_directory_inside_a_checkout_before_the_browser(qbo_env, capsys, monkeypatch):
    import webbrowser

    repo = qbo_env.root / "repo"
    (repo / ".git").mkdir(parents=True)
    qbo_env.set(QBO_REDIRECT_URI=f"http://localhost:{_free_port()}/callback")
    monkeypatch.setenv("LEDGERLENS_TOKEN_DIR", str(repo / "tok"))
    opened = []
    monkeypatch.setattr(webbrowser, "open", opened.append)
    assert main(["qbo-auth", "--timeout", "0.3"]) == 2
    assert "inside the git repository" in capsys.readouterr().out
    assert opened == [] and not (repo / "tok").exists()


def test_a_pull_without_tokens_creates_no_token_directory(qbo_env, capsys):
    qbo_env.set(QBO_REALM_ID="4620816365")
    assert main(["pull-qbo", "--start", "2025-10-01", "--end", "2025-12-31"]) == 2
    assert "run `ledgerlens qbo-auth` first" in capsys.readouterr().out
    assert not qbo_env.tokens.exists()


# --- labels are checked against the ledger before any number is printed ------


@pytest.fixture(scope="module")
def two_ledgers(tmp_path_factory):
    """Two generated ledgers of different lengths: neither one's labels describe the other."""
    root = tmp_path_factory.mktemp("pairs")
    main(["generate", "--start", "2024-01-01", "--end", "2024-03-31", "--out-dir", str(root / "a")])
    main(["generate", "--start", "2024-01-01", "--end", "2024-02-29", "--out-dir", str(root / "b")])
    return root


@pytest.mark.parametrize("command", [
    ["test"], ["score"], ["report", "--no-model"],
])
def test_labels_from_another_ledger_are_refused_before_anything_is_scored(
        two_ledgers, tmp_path, capsys, command):
    out = tmp_path / ("wp.xlsx" if command[0] == "report" else "scored.csv")
    capsys.readouterr()
    code = main([*command, str(two_ledgers / "a" / "ledger.csv"),
                 "--labels", str(two_ledgers / "b" / "labels.csv"), "--out", str(out)])
    printed = capsys.readouterr().out
    assert code == 2, printed
    assert printed.startswith("refused: ") and "have no label" in printed
    assert "Precision" not in printed and "precision" not in printed
    assert not out.exists()
    # the ledger's own labels still pass
    assert main([*command, str(two_ledgers / "a" / "ledger.csv"),
                 "--labels", str(two_ledgers / "a" / "labels.csv"), "--out", str(out)]) == 0


def test_an_unreadable_label_file_is_a_message_not_a_traceback(two_ledgers, tmp_path, capsys):
    ledger = str(two_ledgers / "a" / "ledger.csv")
    capsys.readouterr()
    assert main(["test", ledger, "--labels", str(tmp_path / "missing.csv")]) == 2
    assert capsys.readouterr().out.startswith("error: cannot read labels from")
    columnless = tmp_path / "columnless.csv"
    columnless.write_text("entry_id,anomaly_type\nJE-2024-000001,\n")
    assert main(["test", ledger, "--labels", str(columnless)]) == 2
    assert "missing column 'is_anomaly'" in capsys.readouterr().out
    empty = tmp_path / "empty.csv"
    empty.write_text("")
    assert main(["test", ledger, "--labels", str(empty)]) == 2
    assert capsys.readouterr().out.startswith("error: cannot read labels from")
    guessed = tmp_path / "guessed.csv"
    guessed.write_text("entry_id,is_anomaly,anomaly_type\nJE-2024-000001,,\n")
    assert main(["test", ledger, "--labels", str(guessed)]) == 2
    assert "is_anomaly must be true/false or 1/0" in capsys.readouterr().out


def test_labels_that_mark_no_anomaly_are_scored_by_every_command(two_ledgers, tmp_path, capsys):
    """`score` raised KeyError('mean_model_score') on them while `test` and `summary` ran."""
    ledger, labels = _pair(two_ledgers)
    none = tmp_path / "none.csv"
    none.write_text(re.sub(r"(?m)^([^,\n]+),True,[^\n]*$", r"\1,False,", Path(labels).read_text()))
    assert "True" not in none.read_text()
    for command in (["test"], ["score"], ["summary"]):
        capsys.readouterr()
        assert main([*command, ledger, "--labels", str(none)]) == 0, command
    page = capsys.readouterr().out
    assert "undefined (the labels mark no anomaly)" in page
    assert "Recall by archetype: the labels mark no anomaly" in page
    assert main(["summary", ledger, "--labels", str(none), "--format", "markdown"]) == 0
    markdown = capsys.readouterr().out
    assert "there is no recall by archetype" in markdown and "| Archetype |" not in markdown


# --- summary -----------------------------------------------------------------


def _pair(two_ledgers, name="a"):
    return str(two_ledgers / name / "ledger.csv"), str(two_ledgers / name / "labels.csv")


def test_summary_prints_one_page_whose_numbers_are_the_test_commands(two_ledgers, capsys):
    ledger, labels = _pair(two_ledgers)
    capsys.readouterr()
    assert main(["test", ledger, "--labels", labels]) == 0
    tested = capsys.readouterr().out
    assert main(["summary", ledger, "--labels", labels]) == 0
    page = capsys.readouterr().out
    for label in ("Precision", "Recall", "F1"):
        value = re.search(rf"^{label}\s+(\d\.\d{{3}})$", tested, re.M).group(1)
        assert re.search(rf"^  {label}\s+{re.escape(value)}$", page, re.M), label
    flagged = re.search(r"on ([\d,]+) entrie\(s\)", tested).group(1)
    assert f"  {flagged} of " in page
    for heading in ("Rule tier", "Model tier", "never blended with the rule tier",
                    "Tier agreement", "not a measure of quality",
                    "Detection against the labels (rule tier only)", "Top 10 by rule score", "Notes"):
        assert heading in page
    assert "a question, not a finding" in page and "nine of eleven" in page
    assert page.index("nine of eleven") < page.index("Precision")


@pytest.mark.parametrize("extra", [[], ["--db", "no-such-review.sqlite"]])
def test_summary_json_is_one_document_and_nothing_else_is_printed(
        two_ledgers, tmp_path, monkeypatch, capsys, extra):
    monkeypatch.chdir(tmp_path)
    ledger, labels = _pair(two_ledgers)
    capsys.readouterr()
    assert main(["summary", ledger, "--labels", labels, "--format", "json", "--top", "3", *extra]) == 0
    payload = json.loads(capsys.readouterr().out)  # anything else on stdout would not parse
    assert payload["ledger"] == ledger and len(payload["top_exceptions"]) == 3
    assert payload["detection"]["metrics"]["population"] == payload["entries"]
    assert ("review" in payload) == bool(extra)
    assert not (tmp_path / "no-such-review.sqlite").exists()  # a summary never creates one
    if extra:
        assert payload["review"]["exists"] is False
        assert any("does not exist" in note for note in payload["notes"])


def test_summary_writes_the_rendering_to_out_and_says_so_in_one_line(two_ledgers, tmp_path, capsys):
    ledger, labels = _pair(two_ledgers)
    target = tmp_path / "new" / "dir" / "summary.md"
    capsys.readouterr()
    assert main(["summary", ledger, "--labels", labels, "--format", "markdown",
                 "--out", str(target)]) == 0
    assert capsys.readouterr().out == f"Summary (markdown) written to {target}\n"
    written = target.read_text(encoding="utf-8")
    assert written.startswith("## LedgerLens summary, ") and written.endswith("\n")
    assert "> A flag is a question, not a finding." in written
    assert "### Detection against the labels (rule tier only)" in written
    assert "@" not in written


def test_a_summary_is_only_written_where_git_cannot_pick_it_up(two_ledgers, tmp_path, capsys):
    from ledgerlens.connectors.tokens import repository_root

    if repository_root(tmp_path) is not None:
        pytest.skip(f"{tmp_path} is inside a git repository")
    ledger, _ = _pair(two_ledgers)
    repo = tmp_path / "repo"
    (repo / ".git").mkdir(parents=True)
    (repo / "docs").mkdir()
    (repo / "linked-out").symlink_to(repo / "docs")
    refused = [repo / "summary.md", repo / "docs" / "summary.md", repo / "out",
               repo / "linked-out" / "summary.md", repo / "new" / "out" / "summary.md",
               repo / "out" / ".." / "summary.md"]
    for target in refused:
        capsys.readouterr()
        assert main(["summary", ledger, "--out", str(target)]) == 2, target
        printed = capsys.readouterr().out
        assert printed.startswith("refused: ") and "not under its out/ directory" in printed
    assert sorted(p.name for p in repo.iterdir()) == [".git", "docs", "linked-out"]  # nothing made
    assert list((repo / "docs").iterdir()) == []

    # an out/ that is itself a link to a tracked directory is that directory
    (repo / "out").symlink_to(repo / "docs")
    assert main(["summary", ledger, "--out", str(repo / "out" / "summary.md")]) == 2
    (repo / "out").unlink()

    assert main(["summary", ledger, "--out", str(repo / "out" / "weekly" / "summary.json"),
                 "--format", "json"]) == 0
    assert json.loads((repo / "out" / "weekly" / "summary.json").read_text())["entries"] > 0
    assert main(["summary", ledger, "--out", str(tmp_path / "elsewhere" / "s.txt")]) == 0  # no checkout


def test_summary_refuses_labels_for_another_ledger_and_writes_nothing(
        two_ledgers, tmp_path, capsys, monkeypatch):
    from ledgerlens import cli

    ledger, _ = _pair(two_ledgers, "a")
    _, other_labels = _pair(two_ledgers, "b")
    target = tmp_path / "s.md"
    scored = []
    real = cli.LedgerContext.load
    monkeypatch.setattr(cli.LedgerContext, "load", lambda self: scored.append(1) or real(self))
    capsys.readouterr()
    assert main(["summary", ledger, "--labels", other_labels, "--out", str(target)]) == 2
    printed = capsys.readouterr().out
    assert printed.startswith("refused: ") and "Precision" not in printed
    assert not target.exists()
    assert scored == []  # refused before either tier ran, as test, score and report refuse


@pytest.mark.parametrize("bad", [
    ["--top", "0"], ["--top", "-3"], ["--format", "xml"], ["--out", ""], ["--out", "  "],
])
def test_summary_rejects_arguments_that_name_nothing(two_ledgers, bad):
    ledger, _ = _pair(two_ledgers)
    with pytest.raises(SystemExit) as stopped:
        main(["summary", ledger, *bad])
    assert stopped.value.code == 2


@pytest.mark.filterwarnings("ignore:Could not infer format")
def test_summary_turns_what_it_cannot_read_into_a_message(two_ledgers, tmp_path, capsys):
    import shutil

    ledger, _ = _pair(two_ledgers)
    capsys.readouterr()
    assert main(["summary", str(tmp_path / "missing.csv")]) == 2
    assert capsys.readouterr().out.startswith("error: cannot summarise the ledger")

    not_a_ledger = tmp_path / "notes.csv"
    not_a_ledger.write_text("a,b\n1,2\n")
    assert main(["summary", str(not_a_ledger)]) == 2
    assert capsys.readouterr().out.startswith("error: cannot summarise the ledger")

    # the reason pandas gives quotes the cell it could not read: one line, no control bytes
    text = Path(ledger).read_text()
    first_date = re.search(r"\d{4}-\d{2}-\d{2}", text).group(0)
    (tmp_path / "bad-date.csv").write_text(text.replace(first_date, "\x1b[31mnot-a-date\x07", 1))
    assert main(["summary", str(tmp_path / "bad-date.csv")]) == 2
    printed = capsys.readouterr().out
    assert "not-a-date" in printed and "\x1b" not in printed and "\x07" not in printed
    assert printed.count("\n") == 1

    shutil.copy(ledger, tmp_path / "ledger.csv")
    (tmp_path / "ledger.identity.json").write_text('{"ledger_id": "qbo 123"}')
    assert main(["summary", str(tmp_path / "ledger.csv")]) == 2
    assert "ledger.identity.json" in capsys.readouterr().out

    damaged = tmp_path / "review.sqlite"
    damaged.write_bytes(b"not a database at all")
    assert main(["summary", ledger, "--db", str(damaged)]) == 2
    assert capsys.readouterr().out.startswith("error: ")
    assert damaged.read_bytes() == b"not a database at all"


def test_summary_of_a_quickbooks_pull_measures_nothing_and_says_why(qbo_env, capsys):
    out = qbo_env.root / "data" / "qbo-ledger.csv"
    assert main(["pull-qbo", *QBO_PERIOD, "--fixtures", str(QBO_FIXTURES / "pull"), "--out", str(out)]) == 0
    for fmt in ("text", "markdown", "json"):
        capsys.readouterr()
        assert main(["summary", str(out), "--format", fmt]) == 0
        page = capsys.readouterr().out
        assert "qbo:sandbox-fixtures" in page
        assert "Precision" not in page and "nine of eleven" not in page
        assert "Detection against" not in page and '"detection"' not in page
        assert "No labels were given, so detection quality was not measured" in " ".join(page.split())
        assert "not the generator's default ledger" in " ".join(page.split())
    payload = json.loads(page)
    assert payload["entries"] == 116 and payload["lines"] == 297
    assert payload["default_ledger"] is False and "detection" not in payload


def test_what_the_readme_says_the_tests_make_of_the_recorded_quarter(qbo_env, capsys):
    """The README's paragraph on a pull's flags quotes these figures. They move only with the
    recording or with a test's definition, and then the README has to move with them."""
    import inspect

    import pandas as pd

    from ledgerlens import jets

    out = qbo_env.root / "data" / "qbo-ledger.csv"
    assert _replay_to(out) == 0
    capsys.readouterr()
    assert main(["summary", str(out), "--format", "json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    fired = {test_id: count for test_id, count in payload["flags_by_test"].items() if count}
    assert fired == {"JET-01": 4, "JET-07": 3, "JET-08": 79, "JET-12": 3}
    assert (payload["flagged"], payload["entries"]) == (81, 116)
    assert payload["model_tier"]["flagged"] == 2  # the 2% budget of 116 entries, not a detection
    assert payload["model_tier"]["dropped_constant"] == ["user_freq"]  # the sandbox has one user

    from ledgerlens.ingest import load_csv

    lines = load_csv(out)
    by_jet08 = jets.run_all(lines).query("test_id == 'JET-08'")["entry_id"]
    assert by_jet08.nunique() == len(by_jet08) == 79  # 79 entries, not 79 flags on fewer
    assert int(pd.read_csv(out)["entered_at_estimated"].sum()) == 0  # no keying time was estimated
    per_entry = pd.read_csv(out).groupby("entry_id")
    # JET-12 fires only on manual entries: three flags over three journal entries is all of them
    assert int((per_entry["source"].first() == "Manual").sum()) == 3
    assert int((per_entry.size() > 2).sum()) == 25  # multi-line entries; the generator writes none
    # the definitions the paragraph explains those figures by
    assert (jets.APPROVAL_THRESHOLD, jets.MANUAL_JE_APPROVERS) == (10_000.0, ("controller", "mchen"))
    assert inspect.signature(jets.jet_rare_account_pair).parameters["max_occurrences"].default == 3
    assert inspect.signature(jets.jet_dormant_account).parameters["dormant_days"].default == 120

    # and the sentences themselves, as the README prints them
    readme = " ".join((Path(__file__).resolve().parents[1] / "README.md").read_text().split())
    multi = int((per_entry.size() > 2).sum())
    for said in (
        f"{payload['flagged']} of the {payload['entries']} entries are flagged, {fired['JET-08']} of them by JET-08",
        "All three journal entries fire JET-12",
        "JET-09 looks for an account quiet for more than 120 days",
        f"its 2% budget of {payload['entries']}, not a detection",
        f"more than two lines ({multi} of the {payload['entries']})",
        "an account pairing seen three times or fewer",
        "the generator's $10,000 approval limit",
    ):
        assert said in readme, said


def test_summary_out_is_judged_and_written_as_one_path_tilde_included(
        two_ledgers, tmp_path, monkeypatch, capsys):
    """A shell passes `--out=~/s.md` on as typed. The guard used to judge the home directory
    while the file went to a directory named "~" inside the checkout."""
    from ledgerlens.connectors.tokens import repository_root

    if repository_root(tmp_path) is not None:
        pytest.skip(f"{tmp_path} is inside a git repository")
    ledger, _ = _pair(two_ledgers)
    repo, home = tmp_path / "repo", tmp_path / "home"
    (repo / ".git").mkdir(parents=True)
    home.mkdir()
    monkeypatch.chdir(repo)
    monkeypatch.setenv("HOME", str(home))
    capsys.readouterr()
    for spelled in ("--out=~/s.md", "--out=./~/s2.md"):
        assert main(["summary", ledger, spelled]) == 0, spelled
    printed = capsys.readouterr().out
    assert str(home / "s.md") in printed and "~" not in printed  # it says where the file is
    assert (home / "s.md").exists() and (home / "s2.md").exists()
    assert sorted(p.name for p in repo.iterdir()) == [".git"]  # no directory named "~"

    monkeypatch.setenv("HOME", str(repo))  # a home that is itself the checkout
    assert main(["summary", ledger, "--out=~/s.md"]) == 2
    assert capsys.readouterr().out.startswith("refused: ")
    assert sorted(p.name for p in repo.iterdir()) == [".git"]
    assert main(["summary", ledger, "--out=~/out/s.md"]) == 0
    assert (repo / "out" / "s.md").exists()
    assert main(["summary", ledger, "--out=~no-such-user-7731/s.md"]) == 2
    assert capsys.readouterr().out.splitlines()[-1].startswith("error: cannot write the summary to")


def test_summary_that_cannot_be_written_is_a_message_before_or_after_scoring(
        two_ledgers, tmp_path, monkeypatch, capsys):
    import os

    from ledgerlens import cli

    ledger, _ = _pair(two_ledgers)
    (tmp_path / "a-directory").mkdir()
    (tmp_path / "a-file").write_text("x")
    scored = []
    real = cli.LedgerContext.load
    monkeypatch.setattr(cli.LedgerContext, "load", lambda self: scored.append(1) or real(self))
    capsys.readouterr()
    assert main(["summary", ledger, "--out", str(tmp_path / "a-directory")]) == 2
    assert "it is a directory" in capsys.readouterr().out and scored == []  # said before any scoring
    checkout = tmp_path / "checkout"
    (checkout / ".git").mkdir(parents=True)
    (checkout / "out").mkdir()
    assert main(["summary", ledger, "--out", str(checkout / "out")]) == 2  # the out/ directory itself
    assert "it is a directory" in capsys.readouterr().out
    assert main(["summary", ledger, "--out", str(tmp_path / "a-file" / "s.md")]) == 2
    assert capsys.readouterr().out.startswith("error: cannot write the summary to")
    if os.geteuid() != 0:  # root writes anywhere
        locked = tmp_path / "locked"
        locked.mkdir()
        locked.chmod(0o555)
        try:
            assert main(["summary", ledger, "--out", str(locked / "s.md")]) == 2
            assert capsys.readouterr().out.startswith("error: cannot write the summary to")
        finally:
            locked.chmod(0o755)


def test_test_and_score_print_the_caveat_ahead_of_the_numbers_it_qualifies(two_ledgers, capsys):
    """Their output is a public CI log on every push and every Monday."""
    from ledgerlens import evaluate

    ledger, labels = _pair(two_ledgers)
    capsys.readouterr()
    assert main(["test", ledger, "--labels", labels]) == 0
    printed = " ".join(capsys.readouterr().out.split())
    assert evaluate.DETECTION_CAVEAT in printed
    assert printed.index(evaluate.DETECTION_CAVEAT) < printed.index("Precision")
    assert main(["score", ledger, "--labels", labels]) == 0
    printed = " ".join(capsys.readouterr().out.split())
    assert evaluate.MODEL_TIER_CAVEAT in printed
    assert printed.index(evaluate.MODEL_TIER_CAVEAT) < printed.index("share_of_all_anomalies")
    # without labels there is no measured number, and so no caveat to print
    assert main(["test", ledger]) == 0 and main(["score", ledger]) == 0
    assert "nine of eleven" not in capsys.readouterr().out
