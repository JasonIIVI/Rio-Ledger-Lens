import time

import openpyxl

from ledgerlens import evaluate, jets
from ledgerlens.benford import segmented_benford
from ledgerlens.model import combine, score_ledger
from ledgerlens.report import build_workpaper


def _workpaper(ledger, labels, tmp_path, **kw):
    flags = jets.run_all(ledger)
    scored = jets.score_entries(ledger, flags)
    scores, report = score_ledger(ledger)
    combined = combine(scored, scores)
    metrics = evaluate.evaluate(flags, labels, ledger["entry_id"].unique())
    path = build_workpaper(
        combined, flags, tmp_path / "wp.xlsx",
        benford=segmented_benford(ledger, by="account_code"),
        metrics=metrics, model_report=report.describe(), **kw)
    return path, flags


def test_workpaper_has_expected_sheets(ledger, labels, tmp_path):
    path, _ = _workpaper(ledger, labels, tmp_path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Summary", "Exceptions", "All flags", "Benford", "Methodology"]


def test_summary_states_the_caveat(ledger, labels, tmp_path):
    path, _ = _workpaper(ledger, labels, tmp_path)
    wb = openpyxl.load_workbook(path)
    text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "a question, not a finding" in text


def test_exceptions_sheet_only_contains_flagged_entries(ledger, labels, tmp_path):
    path, flags = _workpaper(ledger, labels, tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Exceptions"]
    header = [c.value for c in ws[1]]
    col = header.index("entry_id") + 1
    ids = {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
    assert ids <= set(flags["entry_id"])


def test_top_n_limits_the_exception_tab(ledger, labels, tmp_path):
    path, _ = _workpaper(ledger, labels, tmp_path, top_n=5)
    wb = openpyxl.load_workbook(path)
    assert wb["Exceptions"].max_row <= 6  # header + 5


def test_methodology_records_the_model_fit(ledger, labels, tmp_path):
    path, _ = _workpaper(ledger, labels, tmp_path)
    wb = openpyxl.load_workbook(path)
    text = " ".join(
        str(c.value) for row in wb["Methodology"].iter_rows() for c in row if c.value
    )
    assert "Isolation Forest" in text
    assert "not an audit procedure" in text


def _note(summary, confidence="medium"):
    return {"summary": summary, "why_flagged": "w", "evidence_to_request": ["x"],
            "suggested_control": "c", "confidence": confidence}


def test_review_columns_show_the_note_the_reviewer_saw(ledger, labels, tmp_path):
    from ledgerlens.review import Decision, ReviewStore

    store = ReviewStore(tmp_path / "review.sqlite", "csv:test")
    flags = jets.run_all(ledger)
    scored = jets.score_entries(ledger, flags)
    top, second, third, fourth, fifth = scored[scored["risk_score"] > 0]["entry_id"].iloc[:5]

    # Decided against the first note, which was then rewritten.
    seen = store.save_narrative(top, _note("A round-thousand manual entry."), model="claude-test")
    store.record(Decision(top, "escalate", "ana", "needs a senior", narrative_id=seen))
    store.save_narrative(top, _note("Rewritten after the decision.", "low"), model="claude-test")
    # Narrated twice, never decided: the latest note is the right one to show.
    store.save_narrative(second, _note("First draft."))
    newest = store.save_narrative(second, _note("Latest draft."))
    # Decided before any note existed.
    store.record(Decision(third, "dismiss", "ben", "routine"))
    # Decided with no note, narrated a second later: the reviewer never saw it.
    store.record(Decision(fourth, "accept", "ben", "checked the invoice"))
    time.sleep(1.1)  # the timestamps carry seconds
    later = store.save_narrative(fourth, _note("Written after the decision."))
    # Narrated, then decided without recording the note (how pre-version decisions look).
    existing = store.save_narrative(fifth, _note("Already there."))
    store.record(Decision(fifth, "dismiss", "ana", "fine"))

    path, _ = _workpaper(ledger, labels, tmp_path, store=store)
    wb = openpyxl.load_workbook(path)
    ws = wb["Exceptions"]
    header = [c.value for c in ws[1]]
    for col in ("narrative_id", "narrative", "narrative_confidence", "narrative_superseded",
                "narrative_seen_by_reviewer", "decision", "reviewer", "note"):
        assert col in header
    rows = {r[header.index("entry_id")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    col = header.index

    assert rows[top][col("decision")] == "escalate"
    assert rows[top][col("narrative_id")] == seen
    assert rows[top][col("narrative")] == "A round-thousand manual entry."
    assert rows[top][col("narrative_confidence")] == "medium"
    assert rows[top][col("narrative_superseded")] is True
    assert rows[top][col("narrative_seen_by_reviewer")] == "yes"

    assert rows[second][col("decision")] is None
    assert rows[second][col("narrative_id")] == newest
    assert rows[second][col("narrative")] == "Latest draft."
    assert rows[second][col("narrative_superseded")] is False
    assert rows[second][col("narrative_seen_by_reviewer")] in (None, "")

    assert rows[third][col("decision")] == "dismiss"
    assert rows[third][col("narrative_id")] is None
    assert rows[third][col("narrative")] is None
    assert rows[third][col("narrative_superseded")] is False
    assert rows[third][col("narrative_seen_by_reviewer")] in (None, "")

    # The later note is shown, but not passed off as the basis of the decision.
    assert rows[fourth][col("narrative_id")] == later
    assert rows[fourth][col("narrative")] == "Written after the decision."
    assert rows[fourth][col("narrative_seen_by_reviewer")] == "no"

    assert rows[fifth][col("narrative_id")] == existing
    assert rows[fifth][col("narrative_seen_by_reviewer")] == "unknown"

    summary = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "Decisions recorded" in summary
