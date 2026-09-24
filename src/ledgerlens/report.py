"""Excel workpaper export.

The output is meant to look like something handed to a senior for review, not
like a data dump: a summary that states what was run and what was found, an
exception tab ordered the way the work should be done, and a methodology tab so
the numbers can be challenged.

Deliberately plain formatting. A workpaper is evidence, not a dashboard.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .review import ReviewStore

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")
LABEL_FONT = Font(bold=True, size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SEVERITY_FILL = {
    "High": PatternFill("solid", fgColor="F8D7DA"),
    "Medium": PatternFill("solid", fgColor="FFF3CD"),
    "Low": PatternFill("solid", fgColor="E2E3E5"),
}


def _autofit(ws, max_width: int = 60) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(11, longest + 2), max_width)


def _write_table(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for j, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=str(name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, name in enumerate(df.columns, start=1):
            value = row[name]
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(name in ("reasons", "reason")))

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


REVIEW_COLUMNS = (
    "narrative_id", "narrative", "narrative_confidence", "narrative_superseded",
    "decision", "reviewer", "note", "decided_at",
)


def _attach_review(exceptions: pd.DataFrame, store: ReviewStore) -> pd.DataFrame:
    """Add what the human loop knows: the narrative the reviewer saw and the latest decision.

    Left merges, so an entry nobody has looked at simply has blank cells - the
    workpaper then doubles as the list of what is still outstanding.

    The narrative shown is the version the decision recorded, because that is
    the text the decision was based on. An entry with no decision, or one
    decided before narratives were versioned, shows the latest version;
    ``narrative_superseded`` says when a newer version exists than the one
    shown, so a reader knows the advice moved on after the reviewer read it.
    """
    decisions = store.current()[
        ["entry_id", "decision", "reviewer", "note", "narrative_id", "decided_at"]
    ]
    latest = store.narratives_frame().set_index("entry_id")["id"]
    versions = store.narratives_frame(latest_only=False).set_index("id")

    out = exceptions.merge(decisions, on="entry_id", how="left")
    newest = pd.to_numeric(out["entry_id"].map(latest)).astype("Int64")
    shown = pd.to_numeric(out["narrative_id"]).astype("Int64").fillna(newest)
    out["narrative_id"] = shown
    out["narrative"] = shown.map(versions["summary"])
    out["narrative_confidence"] = shown.map(versions["confidence"])
    out["narrative_superseded"] = (
        (shown.notna() & (shown != newest)).fillna(False).astype(bool)
    )
    return out[[c for c in out.columns if c not in REVIEW_COLUMNS] + list(REVIEW_COLUMNS)]


def build_workpaper(
    scored: pd.DataFrame,
    flags: pd.DataFrame,
    out_path: str | Path,
    benford: pd.DataFrame | None = None,
    metrics: dict | None = None,
    model_report: str | None = None,
    top_n: int = 250,
    store: ReviewStore | None = None,
) -> Path:
    """Write the exception workpaper and return the path written."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    exceptions = scored[scored["risk_score"] > 0].copy()
    if "model_score" in exceptions.columns:
        exceptions = exceptions.sort_values(
            ["risk_score", "model_score"], ascending=False
        )
    exceptions = exceptions.head(top_n)

    keep = [c for c in (
        "entry_id", "posting_date", "entered_at", "source", "created_by",
        "entry_amount", "risk_score", "model_score", "agreement", "n_flags",
        "tests_fired", "reasons",
    ) if c in exceptions.columns]
    exceptions = exceptions[keep]
    if store is not None:
        exceptions = _attach_review(exceptions, store)
    for col in ("posting_date", "entered_at"):
        if col in exceptions.columns:
            exceptions[col] = pd.to_datetime(exceptions[col]).dt.tz_localize(None)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # placeholder so the sheet exists; real content written below
        pd.DataFrame().to_excel(writer, sheet_name="Summary")
        exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
        flags.to_excel(writer, sheet_name="All flags", index=False)
        if benford is not None and not benford.empty:
            benford.to_excel(writer, sheet_name="Benford", index=False)
        pd.DataFrame().to_excel(writer, sheet_name="Methodology")

        book = writer.book

        # ---- Summary ----
        ws = book["Summary"]
        ws["A1"] = "LedgerLens - Journal Entry Testing Workpaper"
        ws["A1"].font = TITLE_FONT
        rows = [
            ("Prepared", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Population (entries)", int(len(scored))),
            ("Entries flagged", int((scored["risk_score"] > 0).sum())),
            ("Flag rate", "{:.2%}".format(
                (scored["risk_score"] > 0).sum() / max(len(scored), 1))),
            ("Total flags raised", int(len(flags))),
            ("Distinct tests fired", int(flags["test_id"].nunique()) if not flags.empty else 0),
            ("Shown in this workpaper", int(len(exceptions))),
        ]
        if metrics:
            rows += [
                ("", ""),
                ("Evaluated against labels", "yes"),
                ("Precision", metrics.get("precision")),
                ("Recall", metrics.get("recall")),
                ("False positives", metrics.get("false_positives")),
                ("False negatives", metrics.get("false_negatives")),
            ]
        if store is not None:
            summary = store.summary()
            rows += [("", ""), ("Decisions recorded", int(summary["entries"].sum()))]
            rows += [(f"  {r.decision}", int(r.entries)) for r in summary.itertuples()]
            rows += [("Narratives cached", int(len(store.narrative_ids())))]
        for i, (label, value) in enumerate(rows, start=3):
            ws.cell(row=i, column=1, value=label).font = LABEL_FONT
            ws.cell(row=i, column=2, value=value)

        note_row = 3 + len(rows) + 1
        ws.cell(row=note_row, column=1,
                value="A flag is a question, not a finding.").font = LABEL_FONT
        ws.cell(row=note_row + 1, column=1, value=(
            "Each exception below identifies an entry whose characteristics warrant enquiry. "
            "None of them asserts an error or an irregularity."))
        _autofit(ws)

        # ---- Exceptions: severity shading by highest severity fired ----
        ws = book["Exceptions"]
        if not flags.empty and "entry_id" in exceptions.columns:
            worst = (
                flags.assign(rank=flags["severity"].map({"High": 3, "Medium": 2, "Low": 1}))
                .sort_values("rank", ascending=False)
                .drop_duplicates("entry_id")
                .set_index("entry_id")["severity"]
            )
            for i, entry_id in enumerate(exceptions["entry_id"], start=2):
                fill = SEVERITY_FILL.get(worst.get(entry_id))
                if fill is not None:
                    ws.cell(row=i, column=1).fill = fill
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        ws.freeze_panes = "A2"
        _autofit(ws)

        for name in ("All flags", "Benford"):
            if name in book.sheetnames:
                sheet = book[name]
                for cell in sheet[1]:
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
                sheet.freeze_panes = "A2"
                _autofit(sheet)

        # ---- Methodology ----
        ws = book["Methodology"]
        ws["A1"] = "Methodology and limitations"
        ws["A1"].font = TITLE_FONT
        text = [
            "Tiers",
            "  1. Deterministic journal-entry tests. Explainable rules, each producing a written reason.",
            "  2. Unsupervised scoring (Isolation Forest) over engineered per-entry features.",
            "  The two scores are reported separately. They answer different questions and are not blended.",
            "",
            "Scoring",
            "  Rule score = severity-weighted count of distinct tests fired (High 3, Medium 2, Low 1).",
            "  Model score = Isolation Forest, rescaled to 0-1. Flagged by rank, not by absolute cutoff.",
            "",
            "Benford's Law",
            "  MAD is reported with Nigrini's conformity bands alongside chi-square. Chi-square rejects",
            "  conformity on almost any large population, so MAD is the operative statistic.",
            "  Non-conformity is a pointer, not a finding.",
            "",
            "Limitations",
            "  Thresholds are tuned against synthetic data and are a starting point on a real ledger.",
            "  Synthetic anomalies do not adapt to the controls looking for them.",
            "  This is a demonstration of technique. It is not an audit procedure or professional advice.",
        ]
        if model_report:
            text += ["", "Model fit"] + ["  " + ln for ln in model_report.splitlines()]
        for i, line in enumerate(text, start=3):
            cell = ws.cell(row=i, column=1, value=line)
            if line and not line.startswith("  "):
                cell.font = LABEL_FONT
        ws.column_dimensions["A"].width = 105

    return out_path
